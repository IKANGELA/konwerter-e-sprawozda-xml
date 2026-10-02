"""
Konwerter e-sprawozdań finansowych (XML z KRS) do Excela.

Ten moduł nie korzysta z dysku ani z internetu: dostaje bajty pliku XML
i zwraca bajty pliku .xlsx. Dzięki temu ten sam kod działa:
  - w przeglądarce (Pyodide, strona index.html),
  - na komputerze (cli.py).
"""

import io
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SEKCJE = {"Aktywa", "Pasywa", "RZiSPor", "RZiSKalk"}

# Kolory z portfolio ikangela.pl
FIOLET = "6D3A62"
JASNE_TLO = "F0ECEF"


class BladSprawozdania(Exception):
    """Plik nie jest e-sprawozdaniem, które umiemy odczytać."""


# ---------------------------------------------------------------------------
# Odczyt XML
# ---------------------------------------------------------------------------

def nazwa(e):
    """'{http://...}Aktywa' -> 'Aktywa' (nazwa znacznika bez przestrzeni nazw)."""
    return e.tag.split("}")[-1]


def znajdz(korzen, *nazwy):
    """Pierwszy element o jednej z podanych nazw (lub None)."""
    return next((e for e in korzen.iter() if nazwa(e) in nazwy), None)


def tekst(korzen, nazwa_pola):
    e = znajdz(korzen, nazwa_pola)
    return e.text.strip() if e is not None and e.text else ""


def kwoty(e):
    """
    Kwoty pozycji: KwotaA = rok bieżący, KwotaB = rok poprzedni,
    KwotaB1 = rok poprzedni po przekształceniu (korekty).
    """
    return {nazwa(c): float(c.text) for c in e
            if nazwa(c).startswith("Kwota") and c.text and c.text.strip()}


def kod(sciezka):
    """'Aktywa/Aktywa_B/.../Aktywa_B_III_1_C' -> 'B.III.1.c' (zapis jak w ustawie)."""
    ostatni = sciezka.split("/")[-1]
    if ostatni in SEKCJE:
        return ""
    czesci = ostatni.replace("Aktywa_", "").replace("Pasywa_", "").split("_")
    # podpunkty (a, b, c...) małą literą; na 2. miejscu stoją cyfry rzymskie
    return ".".join(c.lower() if i >= 2 and c.isalpha() else c for i, c in enumerate(czesci))


def pozycje(sekcja, etykiety):
    """
    Wszystkie pozycje sekcji w kolejności ze sprawozdania.
    Sprawozdanie to drzewo (A -> A.II -> A.II.1 -> A.II.1.c), więc funkcja
    przechodzi je rekurencyjnie, zapamiętując poziom zagnieżdżenia.
    """
    wynik = []

    def przejdz(e, sciezka, poziom):
        if nazwa(e).startswith("PozycjaUszczegolawiajaca"):
            # pozycja dodana przez spółkę: nazwa jest w treści pliku, nie w schemacie
            opis = next((c.text for c in e.iter() if nazwa(c).startswith("Nazwa") and c.text), "")
            k = {nazwa(c): float(c.text) for c in e.iter()
                 if nazwa(c).startswith("Kwota") and c.text and c.text.strip()}
            wynik.append({"kod": "", "pozycja": f"(pozycja spółki) {opis.strip()}",
                          "poziom": poziom, "kwoty": k})
            return
        k = kwoty(e)
        if k or sciezka in etykiety:
            wynik.append({"kod": kod(sciezka), "pozycja": etykiety.get(sciezka, nazwa(e)),
                          "poziom": poziom, "kwoty": k})
        for dziecko in e:
            if not nazwa(dziecko).startswith("Kwota"):
                przejdz(dziecko, f"{sciezka}/{nazwa(dziecko)}", poziom + 1)

    przejdz(sekcja, nazwa(sekcja), 0)
    return wynik


def odczytaj(dane, etykiety):
    """Bajty pliku XML -> słownik z danymi sprawozdania."""
    try:
        korzen = ET.fromstring(dane)
    except ET.ParseError:
        raise BladSprawozdania("To nie jest poprawny plik XML. Wybierz e-sprawozdanie "
                               "w formacie .xml lub .xml.xades.")

    bilans = znajdz(korzen, "Bilans")
    rzis = znajdz(korzen, "RZiSPor", "RZiSKalk")
    if bilans is None or rzis is None:
        raise BladSprawozdania(
            "W pliku nie ma bilansu lub rachunku zysków i strat. "
            "Obsługiwane są e-sprawozdania jednostek innych (pełny wzór) w formacie XML.")

    s = {
        "firma": tekst(korzen, "NazwaFirmy"),
        "krs": tekst(korzen, "P_1E"),
        "nip": tekst(korzen, "P_1D"),
        "data_od": tekst(korzen, "DataOd"),
        "data_do": tekst(korzen, "DataDo"),
        "wariant": "porównawczy" if nazwa(rzis) == "RZiSPor" else "kalkulacyjny",
        "aktywa": pozycje(znajdz(bilans, "Aktywa"), etykiety),
        "pasywa": pozycje(znajdz(bilans, "Pasywa"), etykiety),
        "rzis": pozycje(rzis, etykiety),
    }
    s["rok"] = int(s["data_do"][:4]) if s["data_do"][:4].isdigit() else None
    s["kontrole"] = kontrole(s)
    s["podsumowanie"] = podsumowanie(s)
    return s


# ---------------------------------------------------------------------------
# Kontrola poprawności i podsumowanie
# ---------------------------------------------------------------------------

def _kwota(lista, kod_, klucz="KwotaA"):
    return next((p["kwoty"].get(klucz) for p in lista if p["kod"] == kod_), None)


def kontrole(s):
    """Zasady księgowe, które w poprawnym sprawozdaniu zawsze muszą się zgadzać."""
    aktywa, pasywa = _kwota(s["aktywa"], ""), _kwota(s["pasywa"], "")
    zysk_bilans = _kwota(s["pasywa"], "A.VI")
    zysk_rzis = _kwota(s["rzis"], "L" if s["wariant"] == "porównawczy" else "O")

    def test(opis, a, b):
        ok = a is not None and b is not None and abs(a - b) <= 1
        return {"opis": opis, "ok": ok}

    return [
        test("Aktywa razem = pasywa razem", aktywa, pasywa),
        test("Zysk netto w rachunku wyników = zysk netto w bilansie", zysk_rzis, zysk_bilans),
    ]


def podsumowanie(s):
    """Kilka kluczowych liczb do szybkiego podglądu na stronie."""
    zysk = "L" if s["wariant"] == "porównawczy" else "O"
    return {
        "Przychody netto ze sprzedaży": _kwota(s["rzis"], "A"),
        "Zysk (strata) netto": _kwota(s["rzis"], zysk),
        "Aktywa razem": _kwota(s["aktywa"], ""),
        "Kapitał własny": _kwota(s["pasywa"], "A"),
    }


# ---------------------------------------------------------------------------
# Zapis do Excela
# ---------------------------------------------------------------------------

def _dopisz(ws, wartosci):
    """
    Dopisuje wiersz, zapisując każdy tekst jako zwykły tekst.
    Bez tego tekst z pliku XML zaczynający się od '=' (np. nazwa firmy
    '=HYPERLINK(...)') zostałby zapisany jako formuła i wykonany przez Excela.
    """
    ws.append(wartosci)
    for komorka in ws[ws.max_row]:
        if isinstance(komorka.value, str):
            komorka.data_type = "s"


def _arkusz(ws, tytul, wiersze, rok):
    pogrubienie = Font(bold=True)
    uzyj_b1 = any(p["kwoty"].get("KwotaB1") for p in wiersze)

    _dopisz(ws, [tytul])
    ws["A1"].font = Font(bold=True, size=13, color=FIOLET)
    kolumny = ["Kod", "Pozycja", f"{rok} (zł)", f"{rok - 1} (zł)"]
    if uzyj_b1:
        kolumny.append(f"{rok - 1} po przekształceniu (zł)")
    _dopisz(ws, kolumny)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=FIOLET)

    for p in wiersze:
        k = p["kwoty"]
        wiersz = [p["kod"], p["pozycja"], k.get("KwotaA"), k.get("KwotaB")]
        if uzyj_b1:
            wiersz.append(k.get("KwotaB1"))
        _dopisz(ws, wiersz)
        r = ws.max_row
        ws.cell(r, 2).alignment = Alignment(indent=p["poziom"], wrap_text=True)
        if p["poziom"] <= 1:
            for c in ws[r]:
                c.font = pogrubienie
                c.fill = PatternFill("solid", fgColor=JASNE_TLO)
        for kol in range(3, len(wiersz) + 1):
            ws.cell(r, kol).number_format = "#,##0.00"

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 70
    for litera in "CDE":
        ws.column_dimensions[litera].width = 20
    ws.freeze_panes = "C3"


def do_excela(s):
    """Słownik ze sprawozdaniem -> bajty pliku .xlsx."""
    wb = Workbook()
    info = wb.active
    info.title = "Informacje"
    for w in (["Spółka", s["firma"]], ["KRS", s["krs"]], ["NIP", s["nip"]],
              ["Okres", f"{s['data_od']} – {s['data_do']}"],
              ["Wariant RZiS", s["wariant"]],
              *[[("✓ " if k["ok"] else "✗ ") + k["opis"]] for k in s["kontrole"]],
              ["Wygenerowano konwerterem e-sprawozdań (ikangela.pl)"]):
        _dopisz(info, w)
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 18, 60

    rok = s["rok"] or 0
    _arkusz(wb.create_sheet("Bilans - aktywa"), f"{s['firma']} – aktywa", s["aktywa"], rok)
    _arkusz(wb.create_sheet("Bilans - pasywa"), f"{s['firma']} – pasywa", s["pasywa"], rok)
    _arkusz(wb.create_sheet("RZiS"),
            f"{s['firma']} – rachunek zysków i strat ({s['wariant']})", s["rzis"], rok)

    bufor = io.BytesIO()
    wb.save(bufor)
    return bufor.getvalue()


def konwertuj(dane, etykiety):
    """Główna funkcja: bajty XML -> (bajty .xlsx, dane sprawozdania)."""
    s = odczytaj(dane, etykiety)
    return do_excela(s), s
